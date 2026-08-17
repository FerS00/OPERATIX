"""Safe, bounded Excel/CSV inspection for the business API."""

from __future__ import annotations

import csv
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

REQUIRED_COLUMNS = frozenset({"sku", "quantity"})
OPTIONAL_COLUMNS = frozenset({"customer_email", "customer_name", "notes", "idempotency_key"})
HEADER_ALIASES = {
    "codigo": "sku",
    "código": "sku",
    "producto": "sku",
    "cantidad": "quantity",
    "correo": "customer_email",
    "email": "customer_email",
    "cliente": "customer_name",
    "nombre_cliente": "customer_name",
    "notas": "notes",
    "clave_idempotencia": "idempotency_key",
}
MAX_PREVIEW_ROWS = 5_000
MAX_PREVIEW_ERRORS = 100


class ExcelFileError(ValueError):
    """Raised when a workbook cannot be safely inspected."""


def preview_file(path: str | Path, *, max_rows: int = MAX_PREVIEW_ROWS) -> dict[str, Any]:
    """Read only the header and bounded rows without changing application data."""
    file_path = Path(path)
    if not file_path.is_file():
        raise ExcelFileError("El archivo no existe en el almacenamiento")
    if max_rows < 1 or max_rows > MAX_PREVIEW_ROWS:
        raise ValueError("max_rows fuera de rango")

    extension = file_path.suffix.lower()
    if extension in {".csv", ".tsv"}:
        delimiter = "\t" if extension == ".tsv" else ","
        try:
            with file_path.open("r", encoding="utf-8-sig", newline="") as stream:
                rows = csv.reader(stream, delimiter=delimiter)
                return _validate_rows(rows, max_rows=max_rows)
        except (OSError, UnicodeError, csv.Error) as error:
            raise ExcelFileError("No se pudo leer el archivo delimitado") from error

    if extension != ".xlsx":
        raise ExcelFileError("El formato no está soportado; usa .xlsx, .csv o .tsv")

    try:
        workbook = load_workbook(file_path, read_only=True, data_only=False)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        return _validate_rows(rows, max_rows=max_rows)
    except (InvalidFileException, OSError, ValueError, KeyError) as error:
        raise ExcelFileError("No se pudo leer el libro Excel") from error
    finally:
        if "workbook" in locals():
            workbook.close()


def _validate_rows(rows: Iterable[Iterable[Any]], *, max_rows: int) -> dict[str, Any]:
    iterator = iter(rows)
    try:
        raw_headers = next(iterator)
    except StopIteration as error:
        raise ExcelFileError("El archivo está vacío") from error

    headers = [_normalize_header(value) for value in raw_headers]
    if not any(headers):
        raise ExcelFileError("La primera fila debe contener encabezados")
    if len(headers) != len(set(headers)):
        raise ExcelFileError("La primera fila contiene encabezados repetidos")

    missing = sorted(REQUIRED_COLUMNS - set(headers))
    if missing:
        raise ExcelFileError(f"Faltan columnas obligatorias: {', '.join(missing)}")

    errors: list[dict[str, Any]] = []
    total_rows = 0
    valid_rows = 0
    for row_number, raw_row in enumerate(iterator, start=2):
        total_rows += 1
        if total_rows > max_rows:
            errors.append(
                {
                    "row": row_number,
                    "column": "*",
                    "message": f"Se alcanzó el límite de {max_rows} filas para la vista previa",
                }
            )
            break
        values = list(raw_row)
        values.extend([None] * (len(headers) - len(values)))
        values = values[: len(headers)]
        row_errors = _validate_row(dict(zip(headers, values, strict=False)), row_number)
        if row_errors:
            errors.extend(row_errors)
        else:
            valid_rows += 1

    return {
        "headers": headers,
        "total_rows": total_rows,
        "valid_rows": valid_rows,
        "invalid_rows": total_rows - valid_rows,
        "errors": errors[:MAX_PREVIEW_ERRORS],
    }


def _validate_row(row: dict[str, Any], row_number: int) -> list[dict[str, Any]]:
    errors: list[dict[str, Any]] = []
    for column, value in row.items():
        if isinstance(value, str) and value.startswith("="):
            errors.append(
                _error(row_number, column, "Las fórmulas no se aceptan en archivos de entrada")
            )

    sku = _text(row.get("sku"))
    if not sku:
        errors.append(_error(row_number, "sku", "SKU obligatorio"))

    quantity = row.get("quantity")
    try:
        numeric_quantity = int(quantity)
        if (
            isinstance(quantity, bool)
            or numeric_quantity <= 0
            or str(numeric_quantity) != str(quantity).strip()
        ):
            raise ValueError
    except (TypeError, ValueError, AttributeError):
        errors.append(_error(row_number, "quantity", "La cantidad debe ser un entero positivo"))

    if not _text(row.get("customer_email")) and not _text(row.get("customer_name")):
        errors.append(
            _error(row_number, "customer_email", "Se requiere correo o nombre del cliente")
        )
    return errors


def _normalize_header(value: Any) -> str:
    text = _text(value).lower().replace(" ", "_")
    return HEADER_ALIASES.get(text, text)


def _text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _error(row: int, column: str, message: str) -> dict[str, Any]:
    return {"row": row, "column": column, "message": message}
