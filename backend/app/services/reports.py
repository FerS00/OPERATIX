"""Currency-safe sales summaries and workbook exports."""

from __future__ import annotations

from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from backend.app.models.business import Sale
from backend.app.models.files import FileRecord
from backend.app.models.security import User
from backend.app.security.audit import record_audit
from backend.app.services.storage import FileStorage, StoredObject, register_file

MONEY_STEP = Decimal("0.01")
REPORT_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def build_sales_report(session: Session) -> dict[str, Any]:
    """Build a report payload without ever summing different currencies together."""
    sales = list(
        session.scalars(
            select(Sale)
            .options(joinedload(Sale.customer), joinedload(Sale.product))
            .order_by(Sale.created_at, Sale.id)
        )
    )
    by_currency: dict[str, dict[str, Any]] = {}
    rows: list[dict[str, Any]] = []
    total_units = 0
    for sale in sales:
        currency = sale.currency
        currency_metrics = by_currency.setdefault(
            currency,
            {"transactions": 0, "units": 0, "total": Decimal("0.00")},
        )
        currency_metrics["transactions"] += 1
        currency_metrics["units"] += sale.quantity
        currency_metrics["total"] = _money(currency_metrics["total"] + sale.total_amount)
        total_units += sale.quantity
        rows.append(
            {
                "id": sale.id,
                "created_at": sale.created_at.isoformat() if sale.created_at else "",
                "customer": sale.customer.name if sale.customer else sale.customer_id,
                "sku": sale.product.sku if sale.product else sale.product_id,
                "quantity": sale.quantity,
                "unit_price": _money(sale.unit_price),
                "total_amount": _money(sale.total_amount),
                "currency": currency,
            }
        )
    return {
        "transactions": len(sales),
        "units": total_units,
        "by_currency": by_currency,
        "rows": rows,
    }


def public_sales_summary(session: Session) -> dict[str, Any]:
    """Return JSON-safe summary metrics without exposing row-level report data."""
    data = build_sales_report(session)
    return {
        "transactions": data["transactions"],
        "units": data["units"],
        "by_currency": {
            currency: {
                "transactions": metrics["transactions"],
                "units": metrics["units"],
                "total": str(metrics["total"]),
            }
            for currency, metrics in sorted(data["by_currency"].items())
        },
    }


def export_sales_report(
    session: Session,
    actor: User,
    storage: FileStorage,
) -> tuple[FileRecord, dict[str, Any]]:
    """Generate a formula-free workbook and register its external file metadata."""
    data = build_sales_report(session)
    workbook = _workbook_from_report(data)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    stored: StoredObject | None = None
    try:
        stored = storage.store(
            output,
            original_name="sales-report.xlsx",
            content_type=REPORT_MIME,
            purpose="export",
        )
        record = register_file(session, actor, stored)
    except Exception:
        if stored is not None:
            storage.delete(stored.storage_path)
        raise
    record_audit(
        session,
        user_id=actor.id,
        action="report.export",
        tool="reports.sales_export",
        parameters_summary=f"file_id={record.id};transactions={data['transactions']}",
        result="created",
        status="success",
    )
    session.commit()
    return record, data


def _workbook_from_report(data: dict[str, Any]) -> Workbook:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    summary.append(["Métrica", "Valor", "Moneda"])
    summary.append(["Transacciones", data["transactions"], ""])
    summary.append(["Unidades", data["units"], ""])
    for currency, metrics in sorted(data["by_currency"].items()):
        summary.append(["Total", metrics["total"], currency])
        summary.append(["Transacciones", metrics["transactions"], currency])
        summary.append(["Unidades", metrics["units"], currency])

    sales_sheet = workbook.create_sheet("Ventas")
    sales_sheet.append(
        ["ID", "Fecha", "Cliente", "SKU", "Cantidad", "Precio unitario", "Total", "Moneda"]
    )
    for row in data["rows"]:
        sales_sheet.append(
            [
                _excel_safe(row["id"]),
                _excel_safe(row["created_at"]),
                _excel_safe(row["customer"]),
                _excel_safe(row["sku"]),
                row["quantity"],
                row["unit_price"],
                row["total_amount"],
                _excel_safe(row["currency"]),
            ]
        )

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        for column in sheet.columns:
            width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 36)
            sheet.column_dimensions[column[0].column_letter].width = width
    return workbook


def _money(value: Decimal) -> Decimal:
    return Decimal(value).quantize(MONEY_STEP, rounding=ROUND_HALF_UP)


def _excel_safe(value: Any) -> Any:
    """Neutralize formula-like strings before placing user data in a workbook."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def report_file_path(storage: FileStorage, record: FileRecord) -> Path:
    """Resolve a stored report path for download after validating its metadata path."""
    return storage.resolve(record.storage_path)
