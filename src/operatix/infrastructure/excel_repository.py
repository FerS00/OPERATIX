"""Excel implementation of the sales repository."""

from __future__ import annotations

from pathlib import Path
from threading import RLock
from typing import Any

from openpyxl import Workbook, load_workbook

from operatix.domain.sales import SALE_HEADERS, SaleRecord

_EXCEL_LOCK = RLock()


class ExcelSalesRepository:
    """Store transactions in a local .xlsx workbook."""

    def __init__(self, path: Path | str, worksheet: str = "Ventas") -> None:
        self.path = Path(path)
        self.worksheet = worksheet

    @staticmethod
    def _row(sale: SaleRecord) -> list[Any]:
        return [
            str(sale.transaction_id),
            sale.created_at.isoformat(),
            sale.customer_name,
            sale.product,
            sale.quantity,
            float(sale.unit_price),
            float(sale.total_amount),
            sale.currency,
            sale.notes or "",
        ]

    def _open(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if self.path.exists():
            workbook = load_workbook(self.path)
            if self.worksheet not in workbook.sheetnames:
                sheet = workbook.create_sheet(self.worksheet)
                sheet.append(SALE_HEADERS)
            return workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = self.worksheet
        sheet.append(SALE_HEADERS)
        return workbook

    def append(self, sale: SaleRecord) -> bool:
        with _EXCEL_LOCK:
            workbook = self._open()
            sheet = workbook[self.worksheet]
            transaction_id = str(sale.transaction_id)
            exists = any(
                str(cell.value) == transaction_id
                for cell in sheet["A"][1:]
                if cell.value is not None
            )
            if exists:
                workbook.close()
                return False

            sheet.append(self._row(sale))
            workbook.save(self.path)
            workbook.close()
            return True

    def list_all(self) -> list[SaleRecord]:
        with _EXCEL_LOCK:
            if not self.path.exists():
                return []

            workbook = load_workbook(self.path, read_only=True, data_only=True)
            if self.worksheet not in workbook.sheetnames:
                workbook.close()
                return []

            sheet = workbook[self.worksheet]
            rows = list(sheet.iter_rows(values_only=True))
            workbook.close()
            if len(rows) < 2:
                return []

            headers = [str(value) for value in rows[0]]
            return [
                SaleRecord.model_validate(dict(zip(headers, row, strict=False)))
                for row in rows[1:]
                if row and row[0]
            ]
